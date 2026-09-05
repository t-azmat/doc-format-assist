"""Tests for the DOCX exporter.

These assert on the *rendered* document rather than the code path, because
every bug this file guards against was a node type that silently vanished from
the output.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from export import render_docx  # noqa: E402
from styles import get_style, spec_from_dict  # noqa: E402


def doc(*nodes):
    return {"type": "doc", "content": list(nodes)}


def para(text):
    return {"type": "paragraph", "content": [{"type": "text", "text": text}]}


def rendered_text(document):
    return "\n".join(p.text for p in document.paragraphs)


def test_page_size_from_spec_is_applied():
    """page_width_in/page_height_in used to be dropped by spec_from_dict and
    never written to the section, so every export came out US Letter."""
    style = spec_from_dict({"page_width_in": 8.27, "page_height_in": 11.69})

    document = render_docx(doc(para("body")), "Title", style)

    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69


def test_page_size_defaults_to_letter():
    document = render_docx(doc(para("body")), "Title", get_style("ieee"))

    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0


def test_margins_and_columns_are_applied():
    document = render_docx(doc(para("body")), "Title", get_style("ieee"))

    section = document.sections[0]
    assert round(section.top_margin.inches, 3) == 0.75
    assert round(section.left_margin.inches, 3) == 0.625


def test_blockquote_is_rendered():
    node = {"type": "blockquote", "content": [para("quoted claim")]}

    document = render_docx(doc(node), "T", get_style("apa"))

    assert "quoted claim" in rendered_text(document)


def test_code_block_preserves_line_breaks():
    node = {
        "type": "codeBlock",
        "content": [{"type": "text", "text": "def f():\n    return 1"}],
    }

    document = render_docx(doc(node), "T", get_style("apa"))

    text = rendered_text(document)
    assert "def f():" in text
    assert "return 1" in text


def test_hard_break_does_not_drop_following_text():
    node = {
        "type": "paragraph",
        "content": [
            {"type": "text", "text": "line one"},
            {"type": "hardBreak"},
            {"type": "text", "text": "line two"},
        ],
    }

    document = render_docx(doc(node), "T", get_style("apa"))

    text = rendered_text(document)
    assert "line one" in text
    assert "line two" in text


def test_nested_lists_are_rendered_at_depth():
    node = {
        "type": "bulletList",
        "content": [
            {
                "type": "listItem",
                "content": [
                    para("outer"),
                    {
                        "type": "bulletList",
                        "content": [
                            {"type": "listItem", "content": [para("inner")]}
                        ],
                    },
                ],
            }
        ],
    }

    document = render_docx(doc(node), "T", get_style("apa"))

    styles = {p.style.name: p.text for p in document.paragraphs if p.text}
    assert styles.get("List Bullet") == "outer"
    assert styles.get("List Bullet 2") == "inner"


def test_headings_use_the_style_typeface_not_word_defaults():
    """python-docx's template renders Heading 1 as blue Calibri Light."""
    node = {
        "type": "heading",
        "attrs": {"level": 1},
        "content": [{"type": "text", "text": "Introduction"}],
    }

    document = render_docx(doc(node), "T", get_style("ieee"))

    heading = next(p for p in document.paragraphs if p.style.name == "Heading 1")
    run = heading.runs[0]
    assert run.font.name == "Times New Roman"
    assert str(run.font.color.rgb) == "000000"


def test_unknown_node_types_render_their_children():
    node = {"type": "someFutureWrapper", "content": [para("still here")]}

    document = render_docx(doc(node), "T", get_style("apa"))

    assert "still here" in rendered_text(document)


def test_table_cell_blocks_become_separate_paragraphs():
    node = {
        "type": "table",
        "content": [
            {
                "type": "tableRow",
                "content": [
                    {
                        "type": "tableCell",
                        "content": [para("first"), para("second")],
                    }
                ],
            }
        ],
    }

    document = render_docx(doc(node), "T", get_style("apa"))

    cell = document.tables[0].cell(0, 0)
    assert [p.text for p in cell.paragraphs] == ["first", "second"]


def heading(text, level):
    return {
        "type": "heading",
        "attrs": {"level": level},
        "content": [{"type": "text", "text": text}],
    }


def test_heading_levels_are_visually_distinguishable():
    """Size alone cannot carry the hierarchy.

    Every preset sets the deepest heading to the body size — IEEE is 10pt body
    and 10pt level 3 — so without a difference in weight or slant a level-3
    heading renders exactly like a paragraph.
    """
    document = render_docx(
        doc(heading("One", 1), heading("Two", 2), heading("Three", 3)),
        "T",
        get_style("ieee"),
    )

    got = {}
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading") and paragraph.runs:
            run = paragraph.runs[0]
            got[paragraph.style.name] = (bool(run.font.bold), bool(run.font.italic))

    assert got["Heading 1"] == (True, False)
    assert got["Heading 2"] == (True, True)
    # Not bold, so it separates from bold body emphasis as well as from roman.
    assert got["Heading 3"] == (False, True)


def test_section_headings_are_not_all_centred():
    """Only a level-1 heading is centred; subsections sit at the measure."""
    document = render_docx(
        doc(heading("One", 1), heading("Two", 2)), "T", get_style("ieee")
    )

    by_style = {
        p.style.name: p.alignment
        for p in document.paragraphs
        if p.style.name.startswith("Heading")
    }

    assert by_style["Heading 1"] == WD_ALIGN_PARAGRAPH.CENTER
    assert by_style["Heading 2"] == WD_ALIGN_PARAGRAPH.LEFT


@pytest.mark.parametrize("style_id", ["ieee", "apa", "acm"])
def test_every_preset_renders(style_id):
    document = render_docx(doc(para("body")), "Title", get_style(style_id))
    assert "Title" in rendered_text(document)
