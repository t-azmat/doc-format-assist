"""Tests for the front-matter section break.

The bug these guard: columns were set on the document, so a two-column paper put
its *title* inside the left column instead of across the page. IEEE and ACM both
run the title and author block full width above a two-column body, which Word
models as two sections — not one.
"""

import sys
from pathlib import Path

import pytest
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from export import render_docx  # noqa: E402
from styles import get_style  # noqa: E402

TWO_COLUMN = ("ieee", "acm")


def para(text):
    return {"type": "paragraph", "content": [{"type": "text", "text": text}]}


def heading(text, level=1):
    return {
        "type": "heading",
        "attrs": {"level": level},
        "content": [{"type": "text", "text": text}],
    }


def doc(*nodes):
    return {"type": "doc", "content": list(nodes)}


def columns_of(section) -> int:
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        return 1
    return int(cols.get(qn("w:num")) or 1)


# --- the split --------------------------------------------------------------

@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_two_column_style_produces_two_sections(style_id):
    document = render_docx(doc(para("body")), "Title", get_style(style_id))

    assert len(document.sections) == 2


@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_front_matter_spans_the_page(style_id):
    """The actual bug: this was 2, which centred the title in the left column."""
    document = render_docx(doc(para("body")), "Title", get_style(style_id))

    assert columns_of(document.sections[0]) == 1


@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_body_section_carries_the_column_count(style_id):
    style = get_style(style_id)
    document = render_docx(doc(para("body")), "Title", style)

    assert columns_of(document.sections[-1]) == style.columns


@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_the_break_is_continuous(style_id):
    """A page break here would push the body onto page two."""
    document = render_docx(doc(para("body")), "Title", get_style(style_id))

    assert document.sections[-1].start_type == WD_SECTION.CONTINUOUS


def test_single_column_style_stays_one_section():
    """APA has nothing to split; it must not gain a section break."""
    document = render_docx(doc(para("body")), "Title", get_style("apa"))

    assert len(document.sections) == 1
    assert columns_of(document.sections[0]) == 1


# --- what lands on which side of the break ----------------------------------

@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_the_title_terminates_the_front_matter_section(style_id):
    """The section break has to sit on the title paragraph. If it drifts onto a
    body paragraph, that paragraph joins the full-width section."""
    document = render_docx(doc(para("body")), "Title", get_style(style_id))

    first = document.paragraphs[0]
    assert "Title" in first.text
    p_pr = first._p.find(qn("w:pPr"))
    assert p_pr is not None and p_pr.find(qn("w:sectPr")) is not None


@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_no_blank_paragraph_is_introduced(style_id):
    """`add_section` parks the break in an empty paragraph; folding it onto the
    title is what keeps a blank line of body height out of the output."""
    content = doc(heading("Introduction"), para("body"))

    two_col = render_docx(content, "Title", get_style(style_id))
    one_col = render_docx(content, "Title", get_style("apa"))

    assert len(two_col.paragraphs) == len(one_col.paragraphs)
    assert all(p.text.strip() for p in two_col.paragraphs)


@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_body_content_survives_the_split(style_id):
    document = render_docx(
        doc(heading("Introduction"), para("body text"), heading("Method")),
        "Title",
        get_style(style_id),
    )

    rendered = "\n".join(p.text for p in document.paragraphs)
    for expected in ("Title", "Introduction", "body text", "Method"):
        assert expected in rendered


# --- geometry ---------------------------------------------------------------

@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_page_geometry_is_identical_either_side_of_the_break(style_id):
    """The new section is cloned from the first, so page size and margins carry
    over. If they ever stop matching, the body reflows onto a different sheet."""
    style = get_style(style_id)
    document = render_docx(doc(para("body")), "Title", style)

    front, body = document.sections[0], document.sections[-1]
    assert front.page_width == body.page_width
    assert front.page_height == body.page_height
    assert front.left_margin == body.left_margin
    assert front.right_margin == body.right_margin
    assert round(body.page_width.inches, 2) == round(style.page_width_in, 2)
    assert round(body.left_margin.inches, 3) == round(style.margins_in["left"], 3)


# --- regressions ------------------------------------------------------------

@pytest.mark.parametrize("style_id", TWO_COLUMN)
def test_contents_page_still_links_across_the_split(style_id):
    """The TOC is built in the body section; the break must not disturb the
    bookmark/hyperlink pairing."""
    import io
    import re
    import zipfile

    content = doc(
        heading("Contents"),
        para("stale ..... 1"),
        heading("Introduction"),
        para("body"),
        heading("Method"),
    )
    document = render_docx(content, "Title", get_style(style_id))

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as archive:
        xml = archive.read("word/document.xml").decode("utf8")

    bookmarks = set(re.findall(r'w:bookmarkStart[^>]*w:name="([^"]+)"', xml))
    links = re.findall(r'<w:hyperlink w:anchor="([^"]+)"', xml)
    assert links and set(links) <= bookmarks


@pytest.mark.parametrize("style_id", ["ieee", "apa", "acm"])
def test_every_preset_still_renders(style_id):
    document = render_docx(doc(heading("A"), para("b")), "T", get_style(style_id))

    assert "T" in "\n".join(p.text for p in document.paragraphs)
